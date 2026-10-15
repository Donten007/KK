"""客户跟进催办的纯逻辑层：记录解析、超期判定、话术生成、交付物构建。

与 executor（followup_chase.py）分离的原因：这四件事没有 IO、没有图语义，
可以脱离数据库与 LangGraph 单测；executor 只做编排与落库。

禁造数据（本岗位的 P0）：每条催办建议必须引用具体跟进记录
（客户名 + 最后跟进日期 + 来源行号），引用不到的记录不生成话术，
宁可留空也不要编一个看起来合理的客户。
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Protocol

DEFAULT_THRESHOLD_DAYS = 7
# 阈值合法区间：防止 0（全量催办）或 3650（永不催办）这类配错值悄悄通过
MIN_THRESHOLD_DAYS = 1
MAX_THRESHOLD_DAYS = 365


@dataclass
class FollowupRecord:
    """一条可溯源的跟进记录：来自上传表格的某一行。"""

    row_no: int  # 来源行号（xlsx 数据行从 2 计，含表头；CSV 同口径）
    customer: str
    last_followup_date: date | None
    note: str

    def citation(self, input_asset_id: str | None = None) -> str:
        """溯源引用。老板点任何一条催办都能回到原始表格行。"""
        source = f"asset:{input_asset_id}" if input_asset_id else "input"
        return f"{source}#row:{self.row_no}"


@dataclass
class ParseResult:
    """解析产物三分：可判定 / 拿不准 / 不可用。

    - valid：客户名与日期齐全，可直接参与超期判定
    - uncertain：日期缺失或解析失败但客户名可读 —— 拿不准是否超期，
      升级给人（F7「主动升级」机制的载体），绝不替人猜
    - unusable：连客户名都没有，无法催办，只能显式列出
    """

    valid: list[FollowupRecord]
    uncertain: list[FollowupRecord]
    unusable: list[FollowupRecord]


def parse_followup_records(raw_rows: list[Any]) -> ParseResult:
    """把 input_payload 里的原始行解析成三分结果。

    只信字段，不猜：日期解析失败 ≠ 今天，客户名缺失 ≠ 「未知客户」。
    """
    valid: list[FollowupRecord] = []
    uncertain: list[FollowupRecord] = []
    unusable: list[FollowupRecord] = []
    for idx, row in enumerate(raw_rows, start=2):  # 行号含表头，数据从第 2 行起
        row_no = row.get("row") if isinstance(row, dict) else None
        if not isinstance(row_no, int) or row_no < 1:
            row_no = idx
        customer = str(row.get("customer") or "").strip() if isinstance(row, dict) else ""
        raw_date = row.get("last_followup_date") if isinstance(row, dict) else None
        note = str(row.get("note") or "").strip() if isinstance(row, dict) else ""
        if not customer:
            unusable.append(FollowupRecord(
                row_no=row_no, customer="", last_followup_date=None, note=note
            ))
            continue
        parsed = _parse_date(raw_date)
        if parsed is None and raw_date not in (None, ""):
            # 日期填了但解析不了：宁可疑真也不放过，归入拿不准
            uncertain.append(FollowupRecord(
                row_no=row_no, customer=customer, last_followup_date=None, note=note
            ))
            continue
        record = FollowupRecord(row_no=row_no, customer=customer, last_followup_date=parsed, note=note)
        if parsed is None:
            uncertain.append(record)
        else:
            valid.append(record)
    return ParseResult(valid=valid, uncertain=uncertain, unusable=unusable)


def days_since(record: FollowupRecord, today: date) -> int:
    assert record.last_followup_date is not None
    return (today - record.last_followup_date).days


def is_overdue(record: FollowupRecord, today: date, threshold_days: int) -> bool:
    """「超过 N 天未跟进」按字面语义：严格大于 N 天才算超期。

    恰好第 N 天不催 —— 少打扰一天的成本远低于误催一次的信任成本。
    """
    if record.last_followup_date is None:
        return False
    return days_since(record, today) > threshold_days


def is_valid_threshold(value: Any) -> bool:
    """阈值校验：bool 是 int 的子类，必须显式排除（True 会被当成 1 天）。"""
    return (
        isinstance(value, int)
        and not isinstance(value, bool)
        and MIN_THRESHOLD_DAYS <= value <= MAX_THRESHOLD_DAYS
    )


class ScriptWriter(Protocol):
    """把一条超期事实写成催办话术。"""

    async def draft(self, record: FollowupRecord, overdue_days: int) -> str: ...


class NoopScriptWriter:
    """无 BYOK 密钥时的兜底：话术只用记录里的真实事实拼装。

    【需资深】为什么「不编造」比「话术漂亮」重要：催办是要发给真客户的，
    链路上一环编造事实（编一个上次沟通内容）就是帮老板得罪客户。
    没有模型可调时，返回基于真实记录的平实话术比返回华丽的幻觉强。
    """

    async def draft(self, record: FollowupRecord, overdue_days: int) -> str:
        last_date = record.last_followup_date.isoformat() if record.last_followup_date else ""
        parts = [f"{record.customer}已经 {overdue_days} 天没有跟进记录（最后跟进：{last_date}）。"]
        if record.note:
            parts.append(f"上次记录：{record.note}。")
        parts.append("建议今天回访确认进展，并把新的跟进结果记回表格。")
        return "".join(parts)


def build_chase_workbook(
    *,
    overdue: list[tuple[FollowupRecord, int, str]],
    uncertain: list[FollowupRecord],
    unusable: list[FollowupRecord],
    today: date,
    input_asset_id: str | None,
) -> bytes:
    """生成催办清单 xlsx。每行带「来源」列（溯源），未确认单列一页。

    overdue 元组：(记录, 超期天数, 话术)。owner 侧交付物不出现任何技术语汇。
    「待人工确认」「无法解析」两页照列来源行号——宁可显式留白也不悄悄吞掉。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "催办话术"
    ws.append(["客户", "最后跟进日期", "超期天数", "催办话术草稿", "来源（溯源引用）"])
    for record, overdue_days, script in overdue:
        last = record.last_followup_date.isoformat() if record.last_followup_date else ""
        ws.append([
            record.customer,
            last,
            overdue_days,
            script,
            record.citation(input_asset_id),
        ])
    ws2 = wb.create_sheet("待人工确认")
    ws2.append(["客户", "最后跟进日期", "为什么需要你确认", "来源（溯源引用）"])
    for record in uncertain:
        ws2.append([
            record.customer,
            record.last_followup_date.isoformat() if record.last_followup_date else "（缺失）",
            "跟进日期缺失或无法识别，无法判断是否超期",
            record.citation(input_asset_id),
        ])
    ws3 = wb.create_sheet("无法解析")
    ws3.append(["来源行号", "问题"])
    for record in unusable:
        ws3.append([record.row_no, "客户名缺失，无法催办"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_DATE_FORMATS = ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日")


def _parse_date(raw: Any) -> date | None:
    """无依赖解析。只认明确格式，认不出就返回 None（调用方归入拿不准）。"""
    if raw in (None, ""):
        return None
    if isinstance(raw, date):
        return raw
    if isinstance(raw, datetime):
        return raw.date()
    text = str(raw).strip()
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        pass
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


__all__ = [
    "DEFAULT_THRESHOLD_DAYS",
    "FollowupRecord",
    "NoopScriptWriter",
    "ParseResult",
    "ScriptWriter",
    "build_chase_workbook",
    "days_since",
    "is_overdue",
    "is_valid_threshold",
    "parse_followup_records",
]
