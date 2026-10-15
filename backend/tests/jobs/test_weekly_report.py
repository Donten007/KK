"""岗位 #2 周报汇总 executor 测试（每 executor ≥2 条：正常 + G4 显式失败）。

核心断言（禁造数据 + G4 + 溯源）：
1. 无日报时**显式失败**（no_daily_reports），不生成空表冒充完成
2. xlsx 每一行都带溯源引用（citation），引用不到的内容不写
3. unknown tool 显式失败（模板 steps_def 与 executor 工具名不同步时立刻暴露）
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from app.jobs.weekly_report import (
    WEEKLY_REPORT_TEMPLATE,
    DailyReportFact,
    NoopSummarizer,
    WeeklyReportExecutor,
)


@dataclass
class _StepCtx:
    step: Any
    goal_id: uuid.UUID = field(default_factory=uuid.uuid4)
    tenant_id: uuid.UUID = field(default_factory=uuid.uuid4)
    attempt: int = 1


def _make_step(tool_name: str) -> _StepCtx:
    from app.agents.graph import StepSpec

    step = StepSpec(ordinal=1, title="t", tool_name=tool_name, tool_args={})
    return _StepCtx(step=step)


@dataclass
class _FakeGoal:
    id: uuid.UUID
    created_at: datetime


@dataclass
class _FakeStep:
    step_ordinal: int
    title: str


def _fake_session_factory(facts: list[DailyReportFact]):
    """最小 fake 会话：满足 tenant_transaction(work_unit + tenant_scope) 的接口面。

    execute 返回的结果对象同时支持 .all()（goal 列表）与 .scalars()（step 列表），
    覆盖业务查询与 RLS GUC 设置两种路径。
    """

    class _GoalsResult:
        def all(self):
            return [(_FakeGoal(id=f.goal_id, created_at=datetime.combine(
                f.report_date, datetime.min.time(), tzinfo=UTC)), f.agent_name) for f in facts]

    class _StepsResult:
        def scalars(self):
            return self

        def all(self):
            return [_FakeStep(step_ordinal=f.step_ordinal, title=f.headline) for f in facts]

    class _GucResult:
        def scalar_one_or_none(self):
            return ""

    class _FakeSession:
        def in_transaction(self) -> bool:
            return True

        async def commit(self) -> None:
            return None

        async def rollback(self) -> None:
            return None

        async def execute(self, *args, **kwargs):
            del kwargs
            stmt = str(args[0]) if args else ""
            if "set_config" in stmt or "current_setting" in stmt:
                return _GucResult()
            if "TaskStep" in stmt or "task_steps" in stmt:
                return _StepsResult()
            return _GoalsResult()

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

    def _factory():  # sessionmaker 是同步可调用对象，返回异步上下文管理器
        return _FakeSession()

    return _factory


def _fact(goal_id: uuid.UUID, headline: str) -> DailyReportFact:
    return DailyReportFact(
        goal_id=goal_id,
        step_ordinal=3,
        report_date=datetime.now(UTC).date() - timedelta(days=1),
        agent_name="市场专员",
        headline=headline,
    )


# ---- 正常路径 ----


@pytest.mark.asyncio
async def test_collect_then_export_produces_traceable_xlsx() -> None:
    goal_id = uuid.uuid4()
    executor = WeeklyReportExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="市场专员",
        session_factory=_fake_session_factory(
            [_fact(goal_id, "竞品 A 上调了标准版价格 10%")]
        ),
        summarizer=NoopSummarizer(),
    )

    collect = await executor(_make_step("collect_daily_reports"))
    assert collect.status == "succeeded", collect.error_code

    summary = await executor(_make_step("summarize_reports"))
    assert summary.status == "succeeded"

    export = await executor(_make_step("generate_xlsx"))
    assert export.status == "succeeded"
    assert export.output_ref and export.output_ref.endswith(".xlsx")

    # 溯源断言：xlsx 内容必须包含 citation，否则「可溯源」是纸面能力
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(executor._xlsx_content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("日期", "员工", "要点", "来源（溯源引用）")
    citation = rows[1][3]
    assert citation == f"goal:{goal_id}#step:3"
    assert "上调了标准版价格" in rows[1][2]


# ---- G4 显式失败：无日报 ----


@pytest.mark.asyncio
async def test_no_reports_is_explicit_failure_not_empty_success() -> None:
    executor = WeeklyReportExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="市场专员",
        session_factory=_fake_session_factory([]),
        summarizer=NoopSummarizer(),
    )
    outcome = await executor(_make_step("collect_daily_reports"))
    assert outcome.status == "failed"
    assert outcome.error_code == "no_daily_reports"


# ---- G4 显式失败：unknown tool（模板与 executor 不同步）----


@pytest.mark.asyncio
async def test_unknown_tool_fails_explicitly() -> None:
    executor = WeeklyReportExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="市场专员",
        session_factory=_fake_session_factory([]),
        summarizer=NoopSummarizer(),
    )
    outcome = await executor(_make_step("not_a_real_tool"))
    assert outcome.status == "failed"
    assert outcome.error_code == "unknown_tool_not_a_real_tool"


# ---- 模板与 executor 一致性 ----
# steps_def 里的每个 tool_name 都必须有对应 _tool_<name> 方法，
# 否则图执行时 unknown_tool 失败。这是「模板与代码同版本上线」的门禁。


def test_template_steps_def_match_executor_tools() -> None:
    executor_tools = {
        name.removeprefix("_tool_")
        for name in dir(WeeklyReportExecutor)
        if name.startswith("_tool_")
    }
    template_tools = {step["tool_name"] for step in WEEKLY_REPORT_TEMPLATE["steps_def"]}
    assert template_tools <= executor_tools, (
        f"模板 steps_def 含 executor 未实现的工具：{template_tools - executor_tools}"
    )


def test_template_has_required_technical_fields() -> None:
    """SPEC §2.2：缺任一组技术字段则 F6/F2 无法落地。"""
    for field_name in ("required_caps", "optional_caps", "expected_idle_steps",
                  "expected_total_steps", "baseline_minutes_per_run"):
        assert field_name in WEEKLY_REPORT_TEMPLATE, f"模板缺 {field_name}"
    assert WEEKLY_REPORT_TEMPLATE["expected_idle_steps"] >= 1