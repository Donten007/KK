"""岗位 #4 客户跟进催办 executor 测试（正常 / 超期识别 / 升级 waiting_approval）。

核心断言（禁造数据 + G4 + 「主动升级给人」）：
1. 正常路径：只给可确认超期的客户出话术，每行带溯源引用（客户/日期/来源行）
2. 超期识别：恰好第 N 天不催、第 N+1 天才催（严格大于语义）
3. 升级路径：跟进日期拿不准 → SoftPauseSignal（/pending 落库链的上游）；
   老板批准后续跑 → 跳过拿不准、不悄悄吞掉（进「待人工确认」页）
4. G4 显式失败：没有输入记录 / 阈值配错 / unknown tool
"""

from __future__ import annotations

import io
import uuid
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from app.agents.graph import SoftPauseSignal
from app.jobs.followup_chase import FOLLOWUP_CHASE_TEMPLATE, FollowupChaseExecutor

TODAY = datetime.now(UTC).date()


@dataclass
class _StepCtx:
    step: Any
    goal_id: uuid.UUID = field(default_factory=uuid.uuid4)
    tenant_id: uuid.UUID = field(default_factory=uuid.uuid4)
    attempt: int = 1


def _make_step(tool_name: str, ordinal: int = 1, title: str = "t") -> _StepCtx:
    from app.agents.graph import StepSpec

    step = StepSpec(ordinal=ordinal, title=title, tool_name=tool_name, tool_args={})
    return _StepCtx(step=step)


def _fake_session_factory(*, payload: dict[str, Any] | None, resolved_pending: bool):
    """最小 fake 会话：满足 tenant_transaction 接口面 + 两种业务查询。

    execute 按语句文本分派：GUC / goal.input_payload / pending_items。
    """

    class _GucResult:
        def scalar_one_or_none(self):
            return ""

    class _PayloadResult:
        def first(self):
            return (payload,) if payload is not None else None

    class _PendingResult:
        def first(self):
            return (uuid.uuid4(),) if resolved_pending else None

    class _FakeSession:
        def in_transaction(self) -> bool:
            return True

        async def commit(self) -> None:
            return None

        async def rollback(self) -> None:
            return None

        async def execute(self, *args, **kwargs):
            del kwargs
            stmt = str(args[0]) if args else ""
            if "set_config" in stmt or "current_setting" in stmt:
                return _GucResult()
            if "input_payload" in stmt:
                return _PayloadResult()
            if "pending_items" in stmt:
                return _PendingResult()
            raise AssertionError(f"fake session 收到未预期的语句：{stmt[:80]}")

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return False

    def _factory():
        return _FakeSession()

    return _factory


def _payload(records: list[dict[str, Any]], **extra: Any) -> dict[str, Any]:
    return {"input": "催办超期客户", "followup_records": records, **extra}


def _overdue_rows() -> list[dict[str, Any]]:
    return [
        {"row": 2, "customer": "深圳华创", "last_followup_date": str(TODAY - timedelta(days=12)), "note": "报价后等回复"},
        {"row": 3, "customer": "广州立信", "last_followup_date": str(TODAY - timedelta(days=8)), "note": ""},
        {"row": 4, "customer": "东莞恒达", "last_followup_date": str(TODAY - timedelta(days=3)), "note": "刚沟通"},
    ]


async def _run_to_build(executor: FollowupChaseExecutor) -> None:
    for tool, ordinal in [
        ("load_followup_table", 1),
        ("identify_overdue", 2),
        ("escalate_uncertain", 3),
        ("draft_scripts", 4),
        ("build_deliverable", 5),
    ]:
        outcome = await executor(_make_step(tool, ordinal=ordinal, title=tool))
        assert outcome.status == "succeeded", (tool, outcome.error_code)


# ---- 正常路径：只催可确认的，逐行可溯源 ----


@pytest.mark.asyncio
async def test_normal_path_produces_traceable_chase_list() -> None:
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload(_overdue_rows()), resolved_pending=False),
    )
    await _run_to_build(executor)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(executor._xlsx_content))
    ws = wb["催办话术"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("客户", "最后跟进日期", "超期天数", "催办话术草稿", "来源（溯源引用）")
    assert [rows[1][0], rows[2][0]] == ["深圳华创", "广州立信"], "刚沟通的东莞恒达不得出现在催办清单"
    assert rows[1][2] == 12 and rows[2][2] == 8
    assert rows[1][4] == "input#row:2" and rows[2][4] == "input#row:3"
    assert "报价后等回复" in rows[1][3], "话术必须引用记录里的真实备注"


@pytest.mark.asyncio
async def test_citation_uses_asset_id_when_present() -> None:
    asset_id = str(uuid.uuid4())
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(
            payload=_payload(_overdue_rows()[:1], input_asset_id=asset_id), resolved_pending=False
        ),
    )
    await _run_to_build(executor)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(executor._xlsx_content))
    rows = list(wb["催办话术"].iter_rows(values_only=True))
    assert rows[1][4] == f"asset:{asset_id}#row:2"


# ---- 超期识别边界（N 可配，默认 7）----


@pytest.mark.asyncio
async def test_overdue_boundary_exactly_n_days_is_not_chased() -> None:
    rows = [
        {"row": 2, "customer": "恰好七天", "last_followup_date": str(TODAY - timedelta(days=7)), "note": ""},
        {"row": 3, "customer": "超过七天", "last_followup_date": str(TODAY - timedelta(days=8)), "note": ""},
    ]
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload(rows), resolved_pending=False),
    )
    await executor(_make_step("load_followup_table", 1))
    outcome = await executor(_make_step("identify_overdue", 2))
    assert outcome.status == "succeeded"
    assert [r[0].customer for r in executor._overdue] == ["超过七天"]


@pytest.mark.asyncio
async def test_threshold_is_configurable() -> None:
    rows = [{"row": 2, "customer": "十四天前", "last_followup_date": str(TODAY - timedelta(days=14)), "note": ""}]
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload(rows, threshold_days=30), resolved_pending=False),
    )
    await executor(_make_step("load_followup_table", 1))
    await executor(_make_step("identify_overdue", 2))
    assert executor._overdue == [], "阈值 30 天时 14 天前的不算超期"


# ---- 升级 waiting_approval（F7「主动升级给人」机制的载体）----


@pytest.mark.asyncio
async def test_uncertain_records_escalate_to_waiting_approval() -> None:
    rows = [
        *_overdue_rows(),
        {"row": 5, "customer": "北京云图", "last_followup_date": None, "note": "日期没填"},
    ]
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload(rows), resolved_pending=False),
    )
    await executor(_make_step("load_followup_table", 1))
    await executor(_make_step("identify_overdue", 2))
    with pytest.raises(SoftPauseSignal) as exc_info:
        await executor(_make_step("escalate_uncertain", 3, title="拿不准的升级确认"))
    signal = exc_info.value
    assert signal.step_title == "拿不准的升级确认" and signal.step_ordinal == 3
    # 升级消息必须带溯源（客户名 + 来源行），不能只说「有记录拿不准」
    assert "北京云图" in signal.user_message and "5" in signal.user_message


@pytest.mark.asyncio
async def test_after_owner_approval_uncertain_skipped_but_not_hidden() -> None:
    """老板批准后续跑：跳过拿不准的，但「待人工确认」页照列（不悄悄吞掉）。"""
    rows = [
        *_overdue_rows(),
        {"row": 5, "customer": "北京云图", "last_followup_date": None, "note": "日期没填"},
    ]
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload(rows), resolved_pending=True),
    )
    await _run_to_build(executor)

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(executor._xlsx_content))
    chased = [r[0] for r in wb["催办话术"].iter_rows(values_only=True)][1:]
    assert "北京云图" not in chased, "拿不准的记录不生成话术（禁造数据）"
    confirm = list(wb["待人工确认"].iter_rows(values_only=True))
    assert confirm[1][0] == "北京云图" and confirm[1][3] == "input#row:5"


# ---- G4 显式失败 ----


@pytest.mark.asyncio
async def test_no_records_is_explicit_failure_not_empty_success() -> None:
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=_payload([]), resolved_pending=False),
    )
    outcome = await executor(_make_step("load_followup_table", 1))
    assert outcome.status == "failed"
    assert outcome.error_code == "no_followup_records"


@pytest.mark.asyncio
async def test_invalid_threshold_fails_explicitly() -> None:
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(
            payload=_payload(_overdue_rows(), threshold_days=0), resolved_pending=False
        ),
    )
    outcome = await executor(_make_step("load_followup_table", 1))
    assert outcome.status == "failed"
    assert outcome.error_code == "invalid_threshold"


@pytest.mark.asyncio
async def test_unknown_tool_fails_explicitly() -> None:
    executor = FollowupChaseExecutor(
        tenant_id=uuid.uuid4(),
        agent_name="销售助理",
        session_factory=_fake_session_factory(payload=None, resolved_pending=False),
    )
    outcome = await executor(_make_step("not_a_real_tool"))
    assert outcome.status == "failed"
    assert outcome.error_code == "unknown_tool_not_a_real_tool"


# ---- 模板与 executor 一致性 ----


def test_template_steps_def_match_executor_tools() -> None:
    executor_tools = {
        name.removeprefix("_tool_")
        for name in dir(FollowupChaseExecutor)
        if name.startswith("_tool_")
    }
    template_tools = {step["tool_name"] for step in FOLLOWUP_CHASE_TEMPLATE["steps_def"]}
    assert template_tools <= executor_tools, (
        f"模板 steps_def 含 executor 未实现的工具：{template_tools - executor_tools}"
    )


def test_template_has_required_technical_fields() -> None:
    for field_name in ("required_caps", "optional_caps", "expected_idle_steps",
                  "expected_total_steps", "baseline_minutes_per_run"):
        assert field_name in FOLLOWUP_CHASE_TEMPLATE, f"模板缺 {field_name}"
    assert FOLLOWUP_CHASE_TEMPLATE["difficulty"] == "medium"
    assert 5 <= FOLLOWUP_CHASE_TEMPLATE["expected_total_steps"] <= 12, "SPEC §2.4 步骤数 5-12"
